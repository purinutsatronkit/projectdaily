from pathlib import Path
import json
import re
import sys

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / 'tmp' / 'video-presentation'
OUTPUT = ROOT / 'artifacts'
FONT = 'Leelawadee UI'


def fonts(obj, size=None, bold=None, color='000000'):
    obj.font.name = FONT
    if size is not None:
        obj.font.size = Pt(size)
    if bold is not None:
        obj.font.bold = bold
    obj.font.italic = False
    obj.font.color.rgb = RGBColor.from_string(color)
    rpr = obj._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn('w:' + key), FONT)
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        rfonts.attrib.pop(qn('w:' + key), None)
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rpr.append(lang)
    lang.set(qn('w:val'), 'th-TH')
    lang.set(qn('w:eastAsia'), 'th-TH')
    lang.set(qn('w:bidi'), 'th-TH')
    if size is not None:
        sizecs = rpr.find(qn('w:szCs'))
        if sizecs is None:
            sizecs = OxmlElement('w:szCs')
            rpr.append(sizecs)
        sizecs.set(qn('w:val'), str(round(size * 2)))


def base_document():
    doc = Document()
    s = doc.sections[0]
    s.page_width = Inches(8.5)
    s.page_height = Inches(11)
    s.top_margin = Inches(.8)
    s.bottom_margin = Inches(.75)
    s.left_margin = Inches(.9)
    s.right_margin = Inches(.9)
    s.header_distance = Inches(.3)
    s.footer_distance = Inches(.3)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3'):
        fonts(doc.styles[name], 11.5)
        pf = doc.styles[name].paragraph_format
        pf.widow_control = True
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.13
        if name.startswith('Heading'):
            pf.keep_with_next = True
    fonts(doc.styles['Title'], 21, True)
    doc.styles['Title'].paragraph_format.space_after = Pt(9)
    fonts(doc.styles['Subtitle'], 10, False)
    doc.styles['Subtitle'].paragraph_format.space_after = Pt(12)
    fonts(doc.styles['Heading 1'], 15, True)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(13)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
    fonts(doc.styles['Heading 2'], 12.5, True)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
    doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
    fonts(doc.styles['Heading 3'], 11.5, True)
    for style in doc.styles:
        for border in list(style._element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    fonts(run, 9)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    run._r.addnext(fld)
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.subject = 'สรุปความรู้จากวิดีโอและแนวทางประยุกต์ใช้'
    return doc


def add_body(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.widow_control = True
    run = p.add_run(text)
    fonts(run)
    return p


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


def add_table(doc, lines):
    data = [[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
    data = [row for row in data if not all(re.fullmatch(r':?-+:?', c) for c in row)]
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.62, 3.05, 2.03] if data[0][0] == 'โครงการ' else [1.45, 2.55, 2.7]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for key in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + key)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '5')
        el.set(qn('w:color'), 'D9D9D9')
        borders.append(el)
    props.append(borders)
    margin = OxmlElement('w:tblCellMar')
    for key, value in (('top', 95), ('bottom', 95), ('left', 110), ('right', 110)):
        el = OxmlElement('w:' + key)
        el.set(qn('w:w'), str(value))
        el.set(qn('w:type'), 'dxa')
        margin.append(el)
    props.append(margin)
    for ri, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        trpr.append(no_split)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader')
            trpr.append(repeat)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            run = p.add_run(data[ri][ci])
            fonts(run, 10.5, ri == 0, 'FFFFFF' if ri == 0 else '000000')
            shade(cell, '1F4E78' if ri == 0 else ('F2F6FA' if ri % 2 else 'FFFFFF'))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)
    return table


def markdown(doc, text, report=False):
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith('|'):
            block = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                block.append(lines[index].strip())
                index += 1
            add_table(doc, block)
            continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            heading = line[3:]
            p = doc.add_paragraph(heading, 'Heading 1')
            if report and heading == 'บทสรุปแบบขยาย':
                p.paragraph_format.page_break_before = True
                p.paragraph_format.space_before = Pt(0)
        elif line.startswith('### '):
            heading = line[4:]
            heading = heading.replace('ระบบจองแบบ Booking.com', 'ระบบจอง')
            doc.add_paragraph(heading, 'Heading 2')
        elif line.startswith('จัดทำวันที่') or line.startswith('ระยะเวลาประมาณ'):
            add_body(doc, line, 'Subtitle')
        else:
            add_body(doc, line)
        index += 1


def build_report():
    doc = base_document()
    markdown(doc, (SOURCE / 'report-revised.md').read_text(encoding='utf-8-sig'), report=True)
    doc.core_properties.title = 'รายงานสรุปวิดีโอเรื่องระบบอัตโนมัติและการออกแบบซอฟต์แวร์'
    OUTPUT.mkdir(parents=True, exist_ok=True)
    path = OUTPUT / 'video-learning-report.docx'
    doc.save(path)
    print(path)


def build_scripts():
    doc = base_document()
    doc.add_paragraph('สคริปต์พูดประกอบรายงานและสไลด์นำเสนอ', 'Title')
    doc.add_paragraph('ระบบอัตโนมัติด้วย n8n และ AI และการออกแบบซอฟต์แวร์', 'Subtitle')
    doc.add_paragraph('ภาคที่ 1 สคริปต์พูดประกอบรายงาน', 'Heading 1')
    report_text = (SOURCE / 'report-speech.md').read_text(encoding='utf-8-sig')
    report_text = re.sub(r'^# .*\n+', '', report_text)
    markdown(doc, report_text)
    p = doc.add_paragraph('ภาคที่ 2 สคริปต์พูดประกอบสไลด์', 'Heading 1')
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    data = json.loads((SOURCE / 'slide-scripts.json').read_text(encoding='utf-8-sig'))
    if isinstance(data, dict):
        slides = data.get('slides', data.get('scripts'))
    else:
        slides = data
    if not isinstance(slides, list) or len(slides) != 10:
        raise ValueError('Expected ten final slide speech scripts')
    p = add_body(doc, 'ระยะเวลารวมประมาณ 9 นาที 15 วินาที', 'Subtitle')
    p.paragraph_format.keep_with_next = True
    for index, slide in enumerate(slides, 1):
        number = slide.get('slide', slide.get('number', index))
        title = slide.get('title', '')
        script = slide.get('script', slide.get('speech', ''))
        seconds = slide.get('seconds', 0)
        duration = f'{seconds} วินาที' if seconds else ''
        heading = f'สไลด์ที่ {number} {title}'.replace(':', '')
        doc.add_paragraph(heading, 'Heading 2')
        if duration:
            p = add_body(doc, f'เวลาพูดประมาณ {duration}', 'Subtitle')
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(4)
        for text in re.split(r'\n\s*\n', script):
            if text.strip():
                add_body(doc, text.strip())
    doc.core_properties.title = 'สคริปต์พูดประกอบรายงานและสไลด์นำเสนอ'
    path = OUTPUT / 'video-learning-speaking-scripts.docx'
    doc.save(path)
    print(path)


if __name__ == '__main__':
    if sys.argv[1] == 'report':
        build_report()
    elif sys.argv[1] == 'scripts':
        build_scripts()
    else:
        raise SystemExit('report or scripts required')
