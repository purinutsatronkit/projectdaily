from pathlib import Path
import json
import re
from zipfile import ZipFile
from lxml import etree
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / 'tmp' / 'video-presentation'
QA = SOURCE / 'word-render'
N = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def plain(doc):
    paragraphs = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(p.text for p in cell.paragraphs)
    return '\n'.join(paragraphs)


results = {}
for name, pages, folder in [('video-learning-report', 5, 'report'), ('video-learning-speaking-scripts', 4, 'scripts')]:
    file = ROOT / 'artifacts' / f'{name}.docx'
    doc = Document(file)
    text = plain(doc)
    assert not re.search(r'https?://|\[\d+\]|บรรณานุกรม|แผนทดลอง', text)
    with ZipFile(file) as archive:
        styles = etree.fromstring(archive.read('word/styles.xml'))
        document = etree.fromstring(archive.read('word/document.xml'))
        assert not styles.xpath('//w:style[@w:styleId="Title"]//w:pBdr', namespaces=N)
        assert not document.xpath('//w:pBdr', namespaces=N)
    pdf = PdfReader(QA / folder / f'{folder}.pdf')
    assert len(pdf.pages) == pages, (name, len(pdf.pages), pages)
    if folder == 'report':
        assert len(doc.tables) == 2
        summary = [p for p in doc.paragraphs if p.text == 'บทสรุปแบบขยาย']
        assert len(summary) == 1 and summary[0].paragraph_format.page_break_before
        assert 'บทสรุปแบบขยาย' in pdf.pages[4].extract_text()
        source = (SOURCE / 'report-revised.md').read_text(encoding='utf-8-sig')
        expected = [line.strip() for line in source.splitlines() if line.strip() and not line.startswith('#') and not line.startswith('|')]
        for paragraph in expected:
            assert paragraph in text, paragraph[:70]
    else:
        source = json.loads((SOURCE / 'slide-scripts.json').read_text(encoding='utf-8-sig'))
        assert len(source['slides']) == 10 and sum(s['seconds'] for s in source['slides']) == 555
        for slide in source['slides']:
            for paragraph in slide['script'].split('\n\n'):
                assert paragraph in text, paragraph[:70]
        assert '9 นาที 15 วินาที' in text
        assert 'sourceScope' not in text and source['sourceScope'] not in text
    results[name] = {'pages': pages, 'bytes': file.stat().st_size, 'visualReview': 'all final pages inspected and passed', 'word': '16.0', 'qaPdf': str(QA / folder / f'{folder}.pdf')}

(QA / 'qa-results.json').write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(results, ensure_ascii=False))
