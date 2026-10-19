from pathlib import Path
from shutil import copy2
import re
from docx import Document

ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / 'tmp' / 'video-presentation' / 'report-revised.md'
FILE = ROOT / 'artifacts' / 'video-learning-report.docx'
BACKUP = Path(__file__).resolve().parent / 'before-remove-timestamps'
BACKUP.mkdir(parents=True, exist_ok=True)
copy2(FILE, BACKUP / FILE.name)
copy2(SOURCE, BACKUP / SOURCE.name)

doc = Document(FILE)
source = SOURCE.read_text(encoding='utf-8-sig')
edits = {
    5: [(' ระยะเวลาประมาณ 43 นาที', '')],
    8: [('เวลา 01:45 ', '')],
    9: [('เวลา 04:44 เป็นตัวอย่างรับข้อมูล', 'ยกตัวอย่างการรับข้อมูล'), ('และเวลา 07:44 เป็นแชตบอต', 'และแชตบอต')],
    10: [('เวลา 10:21 ', ''), ('ส่วนเวลา 13:06 เป็นงาน', 'รวมถึงงาน')],
    11: [('เวลา 15:46 ', ''), ('เวลา 18:38 ', ''), ('และเวลา 22:01 ', 'และ')],
    12: [('เวลา 24:37 เป็นกระบวนการ', 'ยกตัวอย่างกระบวนการ'), ('และเวลา 27:28 เป็นการ', 'และการ')],
    13: [('เวลา 32:25 ', ''), ('เวลา 34:40 ', ''), ('ส่วนเวลา 38:12 ', 'รวมถึง')],
    22: [(' ระยะเวลาประมาณ 3 ชั่วโมง 17 นาที', '')],
    26: [('ช่วง 00:00 ถึง 23:55 ', ''), ('สไลด์ช่วงประมาณ 09:20 ', 'สไลด์')],
    27: [('ในช่วงประมาณ 09:54 ถึง 11:55 ', ''), ('ต่อมาช่วงประมาณ 13:55 เป็นการกำหนด', 'จากนั้นจึงกำหนด'), ('และช่วงประมาณ 16:55 เป็นสถาปัตยกรรม', 'และวางสถาปัตยกรรม')],
    28: [('สไลด์ช่วงประมาณ 18:28 ', 'สไลด์')],
    30: [('ช่วง 23:55 ถึง 50:52 ', ''), (' ที่เวลา 35:23', ''), ('ที่เวลา 40:04 ', ''), ('ที่เวลา 44:37', '')],
    33: [('ช่วง 50:52 ถึง 1:50:45 ', ''), (' ที่เวลา 1:11:02', ''), ('ที่เวลา 1:22:34 ', ''), ('ที่เวลา 1:33:05 ', '')],
    36: [('ช่วง 1:50:45 ถึง 3:11:14 ', ''), ('ที่เวลา 2:13:10 ', ''), ('ที่เวลา 2:22:37 ', ''), ('ที่เวลา 2:31:44 ', '')],
}

before_times = sum(len(re.findall(r'\d+(?::\d+){1,2}', p.text)) for p in doc.paragraphs)
assert before_times == 36, before_times
for index, changes in edits.items():
    p = doc.paragraphs[index]
    original = p.text
    revised = original
    for old, new in changes:
        assert revised.count(old) == 1, (index, old)
        revised = revised.replace(old, new)
    revised = revised.rstrip()
    assert source.count(original) == 1, (index, original)
    source = source.replace(original, revised)
    runs = [r for r in p.runs if r.text]
    assert len(runs) == 1 and runs[0].text == original, (index, len(runs))
    runs[0].text = revised

assert not re.search(r'\d+:\d+', source)
assert not any(re.search(r'\d+:\d+', p.text) for p in doc.paragraphs)
assert 'ระยะเวลาประมาณ 43 นาที' not in source and 'ระยะเวลาประมาณ 3 ชั่วโมง 17 นาที' not in source
assert 'จัดทำวันที่ 19 ตุลาคม 2569' in source
for retain in ('การบริหารเวลา', 'เวลาตอบกลับ', 'พฤติกรรมเวลาระบบมีปัญหา'):
    assert retain in source, retain
doc.save(FILE)
SOURCE.write_text(source, encoding='utf-8')
print(f'Updated {len(edits)} paragraphs; removed {before_times} clip timestamps and both clip durations')
print(FILE)
