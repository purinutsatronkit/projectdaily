from pathlib import Path
import json
import re

from docx.shared import Pt
from build_documents import base_document, add_body, markdown

ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / 'tmp' / 'video-presentation'
JSON_SOURCE = SOURCE / 'v2' / 'slide-scripts.json'
OUTPUT = ROOT / 'artifacts' / 'video-learning-speaking-scripts-detailed.docx'


def duration_th(seconds):
    minutes, remaining = divmod(seconds, 60)
    parts = []
    if minutes:
        parts.append(f'{minutes} นาที')
    if remaining:
        parts.append(f'{remaining} วินาที')
    return ' '.join(parts) or '0 วินาที'


def build():
    data = json.loads(JSON_SOURCE.read_text(encoding='utf-8-sig'))
    slides = data['slides']
    assert len(slides) == 16, len(slides)
    assert [slide['number'] for slide in slides] == list(range(1, 17))
    assert all(isinstance(s['script'], str) and s['script'].strip() for s in slides)
    assert all(isinstance(s['seconds'], int) and s['seconds'] > 0 for s in slides)
    total = sum(s['seconds'] for s in slides)
    doc = base_document()
    doc.add_paragraph('สคริปต์พูดประกอบรายงานและสไลด์ฉบับขยาย', 'Title')
    doc.add_paragraph('ระบบอัตโนมัติด้วย n8n และ AI และการออกแบบซอฟต์แวร์', 'Subtitle')
    doc.add_paragraph('ภาคที่ 1 สคริปต์พูดประกอบรายงาน', 'Heading 1')
    report_text = (SOURCE / 'report-speech.md').read_text(encoding='utf-8-sig')
    report_text = re.sub(r'^# .*\n+', '', report_text)
    markdown(doc, report_text)
    p = doc.add_paragraph('ภาคที่ 2 สคริปต์พูดประกอบสไลด์ฉบับขยาย', 'Heading 1')
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    p = add_body(doc, f'จำนวน 16 สไลด์ ระยะเวลารวมประมาณ {duration_th(total)}', 'Subtitle')
    p.paragraph_format.keep_with_next = True
    for slide in slides:
        heading = f"สไลด์ที่ {slide['number']} {slide['title']}".replace(':', '')
        doc.add_paragraph(heading, 'Heading 2')
        p = add_body(doc, f"เวลาพูดประมาณ {duration_th(slide['seconds'])}", 'Subtitle')
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(4)
        for paragraph in re.split(r'\n\s*\n', slide['script']):
            if paragraph.strip():
                add_body(doc, paragraph.strip())
    doc.core_properties.title = 'สคริปต์พูดประกอบรายงานและสไลด์ฉบับขยาย'
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f'{OUTPUT}\n16 slides and {total} seconds')


if __name__ == '__main__':
    build()
