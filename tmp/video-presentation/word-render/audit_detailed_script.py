from pathlib import Path
from hashlib import sha256
import json
import re
from zipfile import ZipFile
from lxml import etree
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / 'tmp' / 'video-presentation'
FILE = ROOT / 'artifacts' / 'video-learning-speaking-scripts-detailed.docx'
QA = SOURCE / 'word-render' / 'detailed-scripts'
N = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

doc = Document(FILE)
data = json.loads((SOURCE / 'v2' / 'slide-scripts.json').read_text(encoding='utf-8-sig'))
slides = data['slides']
text = '\n'.join(p.text for p in doc.paragraphs)
assert len(slides) == 16
assert sum(s['seconds'] for s in slides) == 970
assert '16 นาที 10 วินาที' in text
assert len(doc.tables) == 0
assert not re.search(r'https?://|\[\d+\]|sourceScope|onSlide|"diagram"', text)
assert data['sourceScope'] not in text
report_speech = (SOURCE / 'report-speech.md').read_text(encoding='utf-8-sig')
for line in report_speech.splitlines():
    if line.strip() and not line.startswith('#'):
        assert line in text
headings = []
for i, p in enumerate(doc.paragraphs):
    if p.style.name == 'Heading 2':
        headings.append(p.text)
        assert p.paragraph_format.keep_with_next is not False
        duration = doc.paragraphs[i+1]
        assert duration.text.startswith('เวลาพูดประมาณ ')
        assert duration.paragraph_format.keep_with_next
        assert doc.paragraphs[i+2].text.strip()
assert len(headings) == 16
for slide, heading in zip(slides, headings):
    assert heading == f"สไลด์ที่ {slide['number']} {slide['title']}".replace(':', '')
    for paragraph in re.split(r'\n\s*\n', slide['script']):
        if paragraph.strip():
            assert paragraph.strip() in text
with ZipFile(FILE) as archive:
    styles = etree.fromstring(archive.read('word/styles.xml'))
    document = etree.fromstring(archive.read('word/document.xml'))
    assert not styles.xpath('//w:style[@w:styleId="Title"]//w:pBdr', namespaces=N)
    assert not document.xpath('//w:pBdr', namespaces=N)

result = {'file': str(FILE), 'slides': 16, 'slideDurationSeconds': 970, 'slideDuration': '16 นาที 10 วินาที', 'reportSpeech': 'included unchanged', 'sourceJSON': 'all 16 spoken scripts preserved', 'excluded': 'sources, sourceScope, onSlide and diagram drawing instructions', 'headingsAndTiming': 'kept with first speech paragraph', 'bytes': FILE.stat().st_size, 'sha256': sha256(FILE.read_bytes()).hexdigest()}
if (QA / 'scripts.pdf').exists():
    pdf = PdfReader(QA / 'scripts.pdf')
    result['pages'] = len(pdf.pages)
    result['rasterPages'] = len(list(QA.glob('page-*.png')))
QA.mkdir(parents=True, exist_ok=True)
(QA / 'structural-qa.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False))
